import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import re
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    md_path = os.path.join(os.path.dirname(__file__), "..", "docs", "PROJECT_REPORT.md")
    output_docx = os.path.join(os.path.dirname(__file__), "..", "docs", "PROJECT_REPORT.docx")

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = docx.Document()

    # Page Margins: 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    in_code_block = False
    code_lines = []
    code_lang = ""

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')

        # Code block boundary
        if line.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lang = line[3:].strip()
                code_lines = []
            else:
                in_code_block = False
                # Write table code block
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F5F5F5")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                
                run = p.add_run("\n".join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x1F, 0x24, 0x21)

                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Markdown Table Detection
        if line.startswith("|") and i + 1 < len(lines) and lines[i+1].strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 3 and "---" in table_lines[1]:
                headers = [c.strip() for c in table_lines[0].split("|")[1:-1]]
                data_rows = []
                for row_line in table_lines[2:]:
                    row_data = [c.strip() for c in row_line.split("|")[1:-1]]
                    data_rows.append(row_data)

                # Render table
                t = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # Format headers
                for col_idx, h_text in enumerate(headers):
                    cell = t.cell(0, col_idx)
                    set_cell_background(cell, "10271D")
                    set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    run = p.add_run(h_text)
                    run.bold = True
                    run.font.name = 'Calibri'
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

                # Format rows
                for r_idx, r_data in enumerate(data_rows):
                    bg_color = "FAFAF7" if r_idx % 2 == 0 else "FFFFFF"
                    for c_idx in range(len(headers)):
                        val = r_data[c_idx] if c_idx < len(r_data) else ""
                        cell = t.cell(r_idx + 1, c_idx)
                        set_cell_background(cell, bg_color)
                        set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        run = p.add_run(val)
                        run.font.name = 'Calibri'
                        run.font.size = Pt(9.5)
                        if "PASS" in val:
                            run.bold = True
                            run.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
                        elif "FAIL" in val:
                            run.bold = True
                            run.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)

                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # Headings
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(0x10, 0x27, 0x1D) # Forest Green
        elif line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[3:].strip())
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0xC8, 0x82, 0x1D) # Forge Amber
        elif line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(line[4:].strip())
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x15, 0x1A, 0x17)
        elif line.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[5:].strip())
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        elif line.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("_________________________________________________________________________________")
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            run.font.size = Pt(8)
        elif line.startswith("* ") or line.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            # Process bold tags
            text = line[2:].strip()
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            match = re.match(r'^\d+\.\s(.*)$', line)
            text = match.group(1) if match else line
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
        elif line.strip() == "":
            pass
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
        i += 1

    doc.save(output_docx)
    print("Document successfully created at:", output_docx)

if __name__ == "__main__":
    create_report()
